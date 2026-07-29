#!/usr/bin/env python3
"""
Vevox Quiz Creator
Helper script for creating properly formatted Vevox quiz Excel files.
"""

import openpyxl
import sys
from typing import List, Dict, Any

# Column order must match Vevox template exactly
COLUMN_ORDER = [
    'Type', 'Title',
    'Choice1', 'Choice2', 'Choice3', 'Choice4', 'Choice5', 'Choice6', 'Choice7', 'Choice8',
    'CorrectAnswer1', 'CorrectAnswer2', 'CorrectAnswer3', 'CorrectAnswer4',
    'CorrectAnswerExplanation',
    'AllowedSelections', 'Minimum', 'Maximum', 'NumberOfDecimals',
    'LowerLabel', 'UpperLabel'
]


def _sanitize_text(value):
    """Replace Unicode characters that Vevox cannot handle with ASCII equivalents."""
    if not isinstance(value, str):
        return value
    value = value.replace('\u2014', '--')   # em dash
    value = value.replace('\u2013', '-')    # en dash
    value = value.replace('\u2212', '-')    # Unicode minus
    value = value.replace('\u2018', "'")    # left single quote
    value = value.replace('\u2019', "'")    # right single quote
    value = value.replace('\u201c', '"')    # left double quote
    value = value.replace('\u201d', '"')    # right double quote
    return value


def create_vevox_quiz(questions: List[Dict[str, Any]], output_file: str, sheet_name: str = "Sample Template") -> None:
    """
    Create a Vevox-compatible Excel quiz file.

    Uses openpyxl directly (not pandas) to ensure empty cells are truly empty
    rather than inlineStr type, which Vevox rejects.

    Args:
        questions: List of question dictionaries with keys:
            - Type: 'MCQ'
            - Title: Question text
            - Choice1-4: Answer choices
            - CorrectAnswer1: Index of correct answer (1-4)
            - CorrectAnswerExplanation: Explanation text
            - AllowedSelections: 1 for single choice
        output_file: Path to output Excel file
        sheet_name: Excel sheet name (default: "Sample Template")
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Write header row
    for col_idx, header in enumerate(COLUMN_ORDER, 1):
        ws.cell(row=1, column=col_idx, value=header)

    # Write question rows — only write non-None values so empty cells stay
    # truly empty (type 'n') rather than becoming inlineStr
    for row_idx, q in enumerate(questions, 2):
        for col_idx, col_name in enumerate(COLUMN_ORDER, 1):
            value = q.get(col_name)
            if value is not None:
                ws.cell(row=row_idx, column=col_idx, value=_sanitize_text(value))

    wb.save(output_file)

    print(f"Created {output_file}")
    print(f"{len(questions)} questions")
    print(f"Sheet name: {sheet_name}")
    print(f"Ready for Vevox import")


def validate_question(q: Dict[str, Any]) -> List[str]:
    """
    Validate a single question dictionary.

    Returns:
        List of validation error messages (empty if valid)
    """
    errors = []

    # Required fields
    if 'Title' not in q or not q['Title']:
        errors.append("Missing or empty 'Title'")

    # Must have at least 2 choices for MCQ
    choices = [q.get(f'Choice{i}') for i in range(1, 9)]
    non_empty_choices = [c for c in choices if c is not None and c != '']
    if len(non_empty_choices) < 2:
        errors.append(f"Need at least 2 choices, found {len(non_empty_choices)}")

    # Must have a correct answer
    if 'CorrectAnswer1' not in q or q['CorrectAnswer1'] is None:
        errors.append("Missing 'CorrectAnswer1'")
    else:
        # Correct answer must be valid choice number
        correct = q['CorrectAnswer1']
        if correct < 1 or correct > len(non_empty_choices):
            errors.append(f"CorrectAnswer1 ({correct}) out of range (1-{len(non_empty_choices)})")

    # Should have explanation
    if 'CorrectAnswerExplanation' not in q or not q['CorrectAnswerExplanation']:
        errors.append("Missing or empty 'CorrectAnswerExplanation'")

    return errors


def validate_questions(questions: List[Dict[str, Any]]) -> bool:
    """
    Validate all questions in the list.

    Returns:
        True if all valid, False otherwise (prints errors)
    """
    all_valid = True
    for i, q in enumerate(questions, 1):
        errors = validate_question(q)
        if errors:
            all_valid = False
            print(f"✗ Question {i} errors:", file=sys.stderr)
            for error in errors:
                print(f"  - {error}", file=sys.stderr)

    return all_valid


if __name__ == "__main__":
    # Example usage
    example_questions = [
        {
            'Type': 'MCQ',
            'Title': 'What is 2 + 2?',
            'Choice1': '3',
            'Choice2': '4',
            'Choice3': '5',
            'Choice4': '22',
            'CorrectAnswer1': 2.0,
            'CorrectAnswerExplanation': 'The sum of 2 and 2 is 4.',
            'AllowedSelections': 1.0
        }
    ]

    if validate_questions(example_questions):
        create_vevox_quiz(example_questions, "example_quiz.xlsx")
    else:
        print("Validation failed!", file=sys.stderr)
        sys.exit(1)
